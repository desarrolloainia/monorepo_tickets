from collections import defaultdict
from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell

from modules.tickets.application.date_ranges import MADRID
from modules.tickets.domain.entities.ticket import UserSpending


def export_spending(users: list[UserSpending], tickets: list[tuple], desde: date, hasta: date) -> bytes:
    workbook = Workbook(write_only=True)
    summary = workbook.create_sheet("Resumen por empleado")
    detail = workbook.create_sheet("Tickets emitidos")

    def append(sheet, values):
        cells = []
        for value in values:
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, str):
                cell.data_type = "s"
            elif isinstance(value, Decimal):
                cell.number_format = '#,##0.00'
            elif isinstance(value, date):
                cell.number_format = 'dd/mm/yyyy hh:mm' if hasattr(value, 'hour') else 'dd/mm/yyyy'
            cells.append(cell)
        sheet.append(cells)

    for sheet in (summary, detail):
        append(sheet, ["Desde", desde, "Hasta", hasta, "Europe/Madrid"])
        sheet.freeze_panes = "A3"
    append(summary, ["Nombre", "Email", "Cantidad", "Importe"])
    append(detail, ["Empleado", "Solicitud", "Código", "Fecha", "Importe unitario"])
    # Both sheets use the same issued-ticket snapshot, even during concurrent approvals.
    counts = defaultdict(int)
    amounts = defaultdict(lambda: Decimal("0.00"))
    for user_id, name, request_id, code, issued_at, price in tickets:
        counts[user_id] += 1
        amounts[user_id] += price
        append(detail, [name, str(request_id), code,
                        issued_at.astimezone(MADRID).replace(tzinfo=None), price])
    for user in users:
        append(summary, [user.nombre, user.email, counts[user.user_id], amounts[user.user_id]])
    total = sum(amounts.values(), Decimal("0.00"))
    append(summary, ["Total", "", len(tickets), total])
    append(detail, ["Total", "", len(tickets), "", total])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
